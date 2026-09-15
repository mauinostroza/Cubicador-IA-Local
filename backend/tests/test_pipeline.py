import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from cubicador.locator import locate_summary_table
from cubicador.pdf_text import PdfText
from cubicador.pipeline import process_pdf
from cubicador.adapters import parse_number


def make_pdf(path: Path) -> None:
    document = SimpleDocTemplate(str(path), pagesize=A4)
    data = [["Ítem", "Descripción", "Unidad", "Cantidad"], ["1.1", "Hormigón H30", "m3", "125,50"], ["1.2", "Acero A630-420H", "kg", "8.450,00"], ["1.3", "Moldaje de fundación", "m2", "310,25"]]
    table = Table(data, colWidths=[60, 220, 70, 90])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    document.build([Paragraph("PLANO DE FUNDACIONES", None), Spacer(1, 250), Paragraph("TABLA RESUMEN DE CANTIDADES", None), table, Paragraph("Notas: Cantidades referenciales", None)])


class LocatorTests(unittest.TestCase):
    def test_accepts_cubicacion_title_aliases(self):
        for title in ("TABLA DE CUBICACIÓN", "CUADRO DE CUBICACION"):
            with self.subTest(title=title):
                text = PdfText((f"{title}\nÍtem  Descripción  Unidad  Cantidad\n1  Hormigón  m3  12,5",))
                candidate = locate_summary_table(text)
                self.assertIsNotNone(candidate)
                self.assertEqual(candidate.title, title)

    def test_prefers_summary_table(self):
        text = PdfText(("OTRA TABLA\nItem  Cantidad\n1  2\nRESUMEN DE CANTIDADES\nÍtem  Descripción  Unidad  Cantidad\n1  Hormigón  m3  12,5",))
        candidate = locate_summary_table(text)
        self.assertIsNotNone(candidate)
        self.assertEqual(candidate.title, "RESUMEN DE CANTIDADES")

    def test_does_not_truncate_long_text(self):
        raw = "x" * 70000 + "\nRESUMEN DE CANTIDADES\nÍtem  Descripción  Unidad  Cantidad\n1  Acero  kg  2"
        document = PdfText((raw,), raw=raw)
        self.assertGreater(len(document.full_text), 60000)
        self.assertIsNotNone(locate_summary_table(document))

    def test_locale_numbers(self):
        for raw, expected in [("8.450,00", 8450), ("8,450.00", 8450), ("125.50", 125.5), ("125,50", 125.5), ("125", 125)]:
            self.assertEqual(parse_number(raw)[:2], (expected, "parsed"))
        self.assertEqual(parse_number("1.234")[1], "ambiguous")


class PipelineTests(unittest.TestCase):
    def test_pdf_to_json_and_excel(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf, xlsx = root / "plano.pdf", root / "cantidades.xlsx"
            make_pdf(pdf)
            result = process_pdf(pdf, xlsx)
            self.assertTrue(result.tabla_encontrada)
            self.assertEqual(len(result.filas), 3)
            self.assertEqual(result.filas[1].quantities[0].numeric_value, 8450.0)
            workbook = load_workbook(xlsx)
            self.assertEqual(workbook.active.max_row, 4)
            self.assertEqual(workbook.sheetnames, ["Cantidades", "Pendientes", "Trazabilidad"])
            self.assertIn("Acero A630-420H", result.filas[1].evidencia.text)
            self.assertEqual(result.filas[1].quantities[0].original, "8.450,00")

    def test_preserves_empty_page_numbering(self):
        from cubicador.adapters import HeuristicInterpreter
        candidate = locate_summary_table(PdfText(("", "RESUMEN DE CANTIDADES\nÍtem  Descripción  Unidad  Cantidad\n1  Hormigón  m3  2,5")))
        result = HeuristicInterpreter().interpret(candidate, "plano.pdf")
        self.assertEqual(result.filas[0].pagina, 2)

    def test_multiple_quantity_columns_are_not_lost(self):
        from cubicador.adapters import HeuristicInterpreter
        candidate = locate_summary_table(PdfText(("RESUMEN DE CANTIDADES\nItem  Descripcion  m3  kg\n1  Fundacion  12,5  450,0",)))
        result = HeuristicInterpreter().interpret(candidate, "plano.pdf")
        self.assertEqual([q.column for q in result.filas[0].quantities], ["m3", "kg"])
        self.assertEqual([q.numeric_value for q in result.filas[0].quantities], [12.5, 450.0])

    def test_absent_table(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "none.pdf"
            SimpleDocTemplate(str(path)).build([Paragraph("Plano sin tabla", None)])
            result = process_pdf(path)
            self.assertFalse(result.tabla_encontrada)

    def test_warns_when_visible_text_may_be_cad_geometry(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cad.pdf"
            SimpleDocTemplate(str(path)).build([Paragraph("ESCALA FECHA REV.", None)])
            result = process_pdf(path)
            self.assertFalse(result.tabla_encontrada)
            self.assertTrue(result.requiere_revision)
            self.assertTrue(any("geometría CAD" in warning for warning in result.advertencias))

    def test_cli_writes_json(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf, xlsx, json_path = root / "plano.pdf", root / "out.xlsx", root / "out.json"
            make_pdf(pdf)
            completed = subprocess.run([sys.executable, "-m", "cubicador.cli", str(pdf), "--excel", str(xlsx), "--json", str(json_path)], capture_output=True, text=True)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertTrue(xlsx.exists() and json_path.exists())


if __name__ == "__main__":
    unittest.main()
