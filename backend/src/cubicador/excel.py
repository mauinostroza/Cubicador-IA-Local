from pathlib import Path
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .models import ExtractionResult


def export_excel(result: ExtractionResult, target: str | Path, max_bytes: int | None = None) -> Path:
    output = Path(target)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cantidades"
    headers = ["Ítem", "Descripción", "Columna/Unidad", "Cantidad", "Cantidad original", "Estado", "Advertencia", "Tipo", "Página", "Archivo", "Tabla"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in result.filas:
        for q in row.quantities:
            sheet.append([row.item, row.descripcion, q.column, q.numeric_value, q.original, q.parse_status, q.warning, row.tipo_fila, row.pagina, result.archivo, result.titulo_tabla])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (14, 48, 18, 16, 18, 12, 30, 12, 10, 28, 36)
    for column, width in zip(sheet.columns, widths):
        sheet.column_dimensions[column[0].column_letter].width = width
    pending = workbook.create_sheet("Pendientes")
    pending.append(["Requiere revisión", "Advertencia", "Candidatos"])
    warnings = result.advertencias or (["Revisar selección de tabla"] if result.requiere_revision else [])
    for warning in warnings:
        pending.append(["Sí" if result.requiere_revision else "No", warning, " | ".join(result.candidatos)])
    for row in result.filas:
        for quantity in row.quantities:
            if quantity.parse_status != "parsed":
                pending.append(["Sí", f"{row.item}/{quantity.column}: {quantity.warning}", quantity.original])
    trace = workbook.create_sheet("Trazabilidad")
    trace.append(["Ítem", "Página", "Línea inicio", "Línea fin", "Evidencia", "Celdas originales", "SHA-256", "Extractor"])
    for row in result.filas:
        trace.append([row.item, row.pagina, row.evidencia.line_start, row.evidencia.line_end, row.evidencia.text, " | ".join(row.cells_original), result.sha256, result.extractor_version])
    fd, temporary = tempfile.mkstemp(prefix=output.stem, suffix=".xlsx", dir=output.parent)
    os.close(fd)
    try:
        workbook.save(temporary)
        if max_bytes is not None and os.path.getsize(temporary) > max_bytes:
            raise ValueError("El Excel excede el tamaño permitido")
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary): os.unlink(temporary)
    return output
